"""Build docs/Sahi_Daam_AI_Demo.pptx for Phase 6 (tagged metrics only)."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Sahi_Daam_AI_Demo.pptx"


def _bullet(tf, lines: list[str]) -> None:
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0


def main() -> int:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        print("Install python-pptx: pip install python-pptx", file=sys.stderr)
        return 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slides: list[tuple[str, list[str]]] = [
        (
            "Sahi Daam AI",
            [
                "Voice-first mandi advisor — Punjab (AICON'26)",
                "Wholesale reference · city routing · harvest hints",
                "DEPLOYED_URL: (fill after deploy)",
            ],
        ),
        (
            "Problem",
            [
                "Farmers hear one broker quote without a wholesale band",
                "Hard to compare net price at another mandi after transport (ESTIMATE math)",
                "Rain/heat timing — need generic, non-chemical guidance",
                "We do not claim proven income gains",
            ],
        ),
        (
            "Architecture",
            [
                "Parse: Gemini (google-genai) or offline lexicon fallback",
                "Band + verdict: scikit-learn quantiles or recent_median_band",
                "Routing: formula with assumed transport/commission (ESTIMATE)",
                "Advisory: RULE-BASED Open-Meteo + trend — NOT ML",
                "Output: Urdu summary; gTTS speaks summary text only",
            ],
        ),
        (
            "Demo — Very low quote (MEASURED)",
            [
                'Input: "40 maund Tomato Okara quote 2800"',
                "Verdict: Very low · Reference: Lahore",
                "Band low/avg/high: 5400 / 5800 / 6000 Rs/maund",
                "Basis: recent_median_band · as_of 2026-10-09",
                "[Placeholder: Streamlit verdict screenshot]",
            ],
        ),
        (
            "Fail closed (MEASURED)",
            [
                'Input: "50 maund Wheat Lahore quote 4500"',
                "Verdict: Unable to assess — no usable Wheat wholesale band",
                "Wheat/Lahore: 0 days in AMIS cache",
                "[Placeholder: Unable to assess UI screenshot]",
            ],
        ),
        (
            "Evaluation (tagged)",
            [
                "AMIS: 434 rows, 2026-07-13..2026-10-09 — MEASURED",
                "All commodity×city pairs still <90 days — MEASURED",
                "Rolling backtest n_test=0 — MEASURED",
                "Spread: Onion gross gap 3.23%, Tomato 4.07% — ESTIMATE",
                "analyze() p50 127 ms, p95 164 ms (n=20) — MEASURED",
                "[Placeholder: reports/figures chart]",
            ],
        ),
        (
            "AI scope & honesty",
            [
                "DEFAULT_MODEL gemini-3.5-flash — set GEMINI_MODEL via check_env",
                "Live typed Gemini: 0/20 if model_not_found — MEASURED",
                "Offline typed parse: 20/20 — MEASURED (not speech accuracy)",
                "No human voice benchmark yet; gTTS synth = SYNTHETIC-AUDIO",
                "Not claimed: ASR %, MAPE, farmer income proof",
            ],
        ),
        (
            "Team & next steps",
            [
                "Team: (fill names in README table)",
                "Try DEPLOYED_URL · Questions?",
                "Built with Cursor / AI coding assistants",
                "Metrics from reports/ on this repo — not LLM estimates",
            ],
        ),
    ]

    for title, bullets in slides:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.0), Inches(1.0))
        box.text_frame.text = title
        box.text_frame.paragraphs[0].font.size = Pt(32)
        box.text_frame.paragraphs[0].font.bold = True
        body = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(5.5))
        _bullet(body.text_frame, bullets)
        for p in body.text_frame.paragraphs:
            p.font.size = Pt(20)

    prs.save(OUT)
    print(f"Wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
